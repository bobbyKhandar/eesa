"""Importing temp must not merge or overwrite any workbook (issue #110)."""

import importlib
import os
import sys
import tempfile
import unittest
from unittest.mock import patch

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

import temp


class TestExcelMergeImportSafe(unittest.TestCase):
    def test_import_has_no_side_effects(self):
        with patch.object(
            temp.os, "listdir", side_effect=AssertionError("listdir called on import")
        ), patch.object(
            temp.Workbook, "save", side_effect=AssertionError("save called on import")
        ):
            importlib.reload(temp)

    def test_merge_appends_branch_column(self):
        from openpyxl import Workbook, load_workbook

        with tempfile.TemporaryDirectory() as tmp:
            for branch in ("IT", "CSE"):
                wb = Workbook()
                ws = wb.active
                ws.append(("q", "a"))
                ws.append((f"{branch}-q1", f"{branch}-a1"))
                wb.save(os.path.join(tmp, f"{branch}.xlsx"))
            out = os.path.join(tmp, "merged.xlsx")

            with open(os.devnull, "w") as devnull, patch.object(
                sys, "stdout", devnull
            ):
                temp.merge_excels_with_branch(tmp, out)

            rows = list(load_workbook(out).active.iter_rows(values_only=True))
            self.assertEqual(rows[0], ("q", "a", "branch"))
            branches = sorted(r[2] for r in rows[1:])
            self.assertEqual(branches, ["CSE", "IT"])

    def test_cli_accepts_input_and_output_paths(self):
        from openpyxl import Workbook, load_workbook

        with tempfile.TemporaryDirectory() as tmp:
            wb = Workbook()
            wb.active.append(("q",))
            wb.active.append(("q1",))
            wb.save(os.path.join(tmp, "IT.xlsx"))
            out = os.path.join(tmp, "merged.xlsx")

            with open(os.devnull, "w") as devnull, patch.object(
                sys, "stdout", devnull
            ):
                temp.main([tmp, out])

            rows = list(load_workbook(out).active.iter_rows(values_only=True))
            self.assertEqual(rows[0], ("q", "branch"))
            self.assertEqual(rows[1], ("q1", "IT"))


if __name__ == "__main__":
    unittest.main()
